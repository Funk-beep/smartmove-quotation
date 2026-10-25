import openpyxl, json
from pathlib import Path

xlsx_path = Path(r"C:\Users\razor\Desktop\22 10 2026\SMART MOVE CO v8.4.14 ADVANCED SALES & COMMUNICATION SYSTEM.xlsx")
wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

cfg_ws = wb['Company Configuration']
cfg_rows = list(cfg_ws.iter_rows(values_only=True))
cfg_map = {}
for row in cfg_rows:
    if not row or row[0] is None:
        continue
    key = str(row[0]).strip()
    if not key:
        continue
    value = row[1] if len(row) > 1 else ''
    cfg_map[key] = value

svc_ws = wb['AI Service Card Description']
svc_rows = list(svc_ws.iter_rows(values_only=True))
if svc_rows:
    headers = svc_rows[0]
    service_entries = []
    for row in svc_rows[1:]:
        if not row or row[0] is None:
            continue
        obj = {headers[i]: row[i] if i < len(row) else '' for i in range(len(headers))}
        if str(obj.get('Active', '')).strip().upper() == 'TRUE':
            service_entries.append(obj)
else:
    service_entries = []

payload = {
    'company': cfg_map,
    'services': service_entries
}
print(json.dumps(payload, ensure_ascii=False, indent=2)[:60000])
