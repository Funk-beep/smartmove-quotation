import html
import json
import re
from pathlib import Path

import openpyxl

WORKBOOK = Path(r"C:\Users\razor\Desktop\22 10 2026\SMART MOVE CO v8.4.14 ADVANCED SALES & COMMUNICATION SYSTEM.xlsx")
OUT = Path(__file__).with_name("landing_page.html")


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


def clean_url(value):
    value = normalize(value)
    if not value:
        return ""
    if value.startswith("http://") or value.startswith("https://"):
        return value
    return "https://" + value


def normalize_image_url(value):
    value = normalize(value)
    if not value:
        return ""

    if "drive.google.com" in value:
        match = re.search(r"/file/d/([A-Za-z0-9_-]+)", value)
        if match:
            return "https://drive.google.com/uc?export=view&id=" + match.group(1)
        match = re.search(r"[?&]id=([A-Za-z0-9_-]+)", value)
        if match:
            return "https://drive.google.com/uc?export=view&id=" + match.group(1)

    return value


def find_local_image_for_title(title):
    title = normalize(title)
    if not title:
        return ""

    normalized_title = re.sub(r"[^a-z0-9]+", " ", title.lower()).strip()
    best_match = ""
    best_score = -1

    for path in Path(__file__).resolve().parent.iterdir():
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".avif"}:
            continue

        candidate = re.sub(r"[^a-z0-9]+", " ", path.stem.lower()).strip()
        candidate = re.sub(r"^smart move social media ad\s*[-_]*\s*", "", candidate)
        if not candidate:
            continue

        if candidate == normalized_title:
            return path.name

        score = 0
        if normalized_title in candidate or candidate in normalized_title:
            score = 2
        if "" in candidate:
            score = 3
        if score > best_score:
            best_score = score
            best_match = path.name

    return best_match


wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)

cfg_rows = list(wb["Company Configuration"].iter_rows(values_only=True))
company = {}
for row in cfg_rows:
    if not row or row[0] is None:
        continue
    key = normalize(row[0])
    if not key:
        continue
    company[key] = row[1] if len(row) > 1 else ""

service_rows = list(wb["AI Service Card Description"].iter_rows(values_only=True))
headers = service_rows[0] if service_rows else []
services = []
seen = set()
for row in service_rows[1:]:
    if not row or row[0] is None:
        continue
    obj = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
    if normalize(obj.get("Active")).upper() != "TRUE":
        continue
    title = normalize(obj.get("Card Title") or obj.get("Service Name"))
    if not title or title in seen:
        continue
    seen.add(title)
    local_image = find_local_image_for_title(title)
    if local_image:
        image_value = local_image
    else:
        image_value = normalize_image_url(obj.get("Image Source") or obj.get("Image URL") or "")
    services.append({
        "title": title,
        "image": image_value,
        "description": normalize(obj.get("Customer-Facing Description") or ""),
        "group": normalize(obj.get("Service Group") or "")
    })

company_name = normalize(company.get("Company_Name") or "Smart Move Co")
if "{{COMPANY_NAME}}" in company_name:
    company_name = "Smart Move Co"
company_compact = normalize(company.get("Company_Name_Compact") or company_name)
logo_url = clean_url(company.get("Company_Logo_Url") or "")
website_domain = normalize(company.get("Website_Domain") or "www.smartmoveco.co.za")
company_slogan = normalize(company.get("Slogan") or company.get("Company_Slogan")
founded_year = normalize(company.get("Founded_Year") or "")
request_form_url = normalize(company.get("Request_Quotation_Form_Url") or "")
if not request_form_url or "docs.google.com" in request_form_url.lower() or "/forms/" in request_form_url.lower():
    request_form_url = "#quote"
elif not request_form_url.startswith(("http://", "https://")):
    request_form_url = "#quote"
image_viewer_url = clean_url(company.get("Social_Image_Viewer_Url") or "")
primary_colour = normalize(company.get("Primary_Brand_Color") or company.get("Primary_Colour") or "#1976D2")
dark_grey_colour = normalize(company.get("Dark_Grey_Brand_Color") or "#424242")
footer_grey_colour = normalize(company.get("Email_Footer_Copyright_Color") or "#888888")
whatsapp_colour = normalize(company.get("Email_WhatsApp_Color") or "#25D366")
alt_text = normalize(company.get("Email_Logo_Alt_Text") or company_name)
if "{{COMPANY_NAME}}" in alt_text:
    alt_text = alt_text.replace("{{COMPANY_NAME}}", company_name)

# Build the service menu with a professional set of cards.
featured_services = services[:6]
service_options = "\n".join(
    f'<option value="{html.escape(item["title"], quote=True)}">{html.escape(item["title"])}</option>'
    for item in services
)

hero_cards = "\n".join(
    f'''
    <div class="mini-card">
      <img src="{html.escape(item['image']) if item['image'] else 'https://images.unsplash.com/photo-1552664730-d307ca884978?auto=format&fit=crop&w=900&q=80'}" alt="{html.escape(item['title'])}" />
      <div>
        <strong>{html.escape(item['title'])}</strong>
      </div>
    </div>
    '''
    for item in featured_services
)

page = f"""<!doctype html>
<html lang=\"en\">
  <head>
    <meta charset=\"UTF-8\" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>{html.escape(company_name)} | Quote Request</title>
    <meta name="description" content="{html.escape(company_slogan)}" />
    <style>
      :root {{
        --primary: {primary_colour};
        --dark-grey: {dark_grey_colour};
        --footer-grey: {footer_grey_colour};
        --whatsapp: {whatsapp_colour};
        --ink: #0f172a;
        --muted: #4b5563;
        --soft: #f8fafc;
        --white: #ffffff;
        --line: #dbe3ee;
      }}
      * {{ box-sizing: border-box; }}
      body {{
        margin: 0; font-family: Arial, Helvetica, sans-serif; color: var(--ink);
        background: linear-gradient(180deg, #f7fafd 0%, #eef5fb 100%);
      }}
      a {{ color: inherit; text-decoration: none; }}
      .container {{ max-width: 1200px; margin: 0 auto; padding: 0 20px; }}
      header {{
        background: rgba(255,255,255,0.9); backdrop-filter: blur(10px);
        border-bottom: 1px solid var(--line); position: sticky; top: 0; z-index: 100;
      }}
      .nav {{ display: flex; align-items: center; justify-content: space-between; padding: 18px 0; }}
      .brand {{ display: flex; align-items: center; gap: 14px; }}
      .brand img {{ width: 52px; height: 52px; object-fit: contain; border-radius: 14px; }}
      .brand strong {{ display: block; font-size: 1.1rem; color: var(--footer-grey); }}
      .brand span {{ color: var(--footer-grey); font-size: 0.8rem; }}
      .nav-links {{ display: flex; gap: 18px; align-items: center; color: var(--muted); font-weight: 600; }}
      .nav-links a:hover {{ color: var(--primary); }}
      .cta-btn, .primary-btn {{
        display: inline-flex; align-items: center; justify-content: center; border: none; border-radius: 999px;
        padding: 14px 22px; font-weight: 700; cursor: pointer; transition: transform .2s ease;
      }}
      .cta-btn:hover, .primary-btn:hover {{ transform: translateY(-1px); }}
      .cta-btn {{ background: var(--whatsapp); color: #ffffff; }}
      .primary-btn {{ background: var(--primary); color: var(--white); }}
      .hero {{ padding: 70px 0 40px; }}
      .hero-wrap {{ display: grid; grid-template-columns: 1.12fr 0.88fr; gap: 46px; align-items: center; }}
      .eyebrow {{
        display: inline-flex; padding: 8px 14px; border-radius: 999px; background: rgba(10,77,122,0.08);
        color: var(--footer-grey); font-size: 0.78rem; font-weight: 800; letter-spacing: .08em; text-transform: uppercase;
      }}
      h1 {{ font-size: clamp(2.2rem, 4vw, 4rem); line-height: 1.08; margin: 18px 0 16px; color: var(--footer-grey); }}
      .lead {{ color: var(--footer-grey); font-size: 1.08rem; line-height: 1.8; max-width: 620px; }}
      .hero-actions {{ display: flex; gap: 16px; margin-top: 28px; flex-wrap: wrap; }}
      .surface {{ background: var(--white); border: 1px solid var(--line); border-radius: 22px; box-shadow: 0 20px 45px rgba(15,23,42,.08); }}
      .hero-panel {{ padding: 28px; }}
      .hero-panel h3 {{ margin: 0 0 18px; font-size: 1.15rem; color: var(--footer-grey); }}
      .stat-grid {{ display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 12px; margin-top: 18px; }}
      .stat {{ background: var(--soft); border-radius: 16px; border: 1px solid var(--line); padding: 18px 14px; text-align: center; }}
      .stat strong {{ display: block; font-size: 1.4rem; color: var(--primary); }}
      .stat span {{ font-size: 0.82rem; color: var(--muted); }}
      .feature-grid {{ padding: 20px 0 60px; }}
      .mini-card-grid {{ display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 20px; margin-top: 30px; }}
      .mini-card {{ background: var(--white); border: 1px solid var(--line); border-radius: 20px; overflow: hidden; }}
      .mini-card img {{ width: 100%; height: 180px; object-fit: cover; display: block; }}
      .mini-card div {{ padding: 16px; }}
      .mini-card strong {{ font-size: 1rem; }}
      .section-label {{ color: var(--primary); font-weight: 800; letter-spacing: .08em; text-transform: uppercase; font-size: .75rem; }}
      .section-header {{ display: flex; justify-content: space-between; align-items: end; gap: 16px; margin-bottom: 18px; }}
      .section-header h2 {{ margin: 8px 0 0; font-size: clamp(1.7rem, 2vw, 2.5rem); color: var(--footer-grey); }}
      .form-shell {{ padding: 70px 0 90px; }}
      .form-wrap {{ display: grid; grid-template-columns: 0.9fr 1.1fr; gap: 28px; align-items: start; }}
      .info-panel {{ padding: 28px; background: linear-gradient(180deg, var(--dark-grey) 0%, var(--dark-grey) 100%); color: var(--white); border-radius: 24px; }}
      .info-panel h3 {{ margin: 0 0 18px; font-size: 1.8rem; color: var(--white); }}
      .info-panel p {{ color: rgba(255,255,255,.85); line-height: 1.8; }}
      .bullet-list {{ list-style: none; padding: 0; margin: 20px 0 0; display: grid; gap: 10px; }}
      .bullet-list li {{ display: flex; align-items: center; gap: 10px; }}
      .bullet-list li::before {{ content: "✓"; display: inline-flex; width: 24px; height: 24px; align-items: center; justify-content: center; border-radius: 50%; background: rgba(255,255,255,.15); font-weight: 700; }}
      form {{ background: var(--white); border: 1px solid var(--line); border-radius: 24px; padding: 30px; }}
      .field {{ display: flex; flex-direction: column; gap: 8px; margin-bottom: 20px; }}
      label {{ font-size: 0.9rem; font-weight: 700; color: var(--ink); }}
      input, select {{ width: 100%; border: 1px solid var(--line); border-radius: 12px; padding: 14px 15px; font-size: 1rem; color: var(--ink); background: #fff; }}
      input:focus, select:focus {{ outline: 2px solid rgba(10,77,122,.15); border-color: var(--primary); }}
      .submit-row {{ margin-top: 8px; display: flex; justify-content: flex-end; }}
      .foot {{ padding: 28px 0 50px; color: var(--muted); text-align: center; }}
      @media (max-width: 900px) {{
        .hero-wrap, .form-wrap {{ grid-template-columns: 1fr; }}
        .mini-card-grid {{ grid-template-columns: 1fr; }}
        .nav-links {{ display: none; }}
      }}
    </style>
  </head>
  <body>
    <header>
      <div class="container nav">
        <div class="brand">
          <img src="{html.escape(logo_url) if logo_url else 'https://images.unsplash.com/photo-1552664730-d307ca884978?auto=format&fit=crop&w=1200&q=80'}" alt="{html.escape(alt_text)}" />
          <div>
            <strong>{html.escape(company_name)}</strong>
            <span>{html.escape(website_domain)}</span>
          </div>
        </div>
        <nav class="nav-links">
          <a href="#services">Services</a>
          <a href="#quote">Request a quote</a>
        </nav>
      </div>
    </header>

    <main>
      <section class="hero">
        <div class="container hero-wrap">
          <div>
            <h1>{html.escape(company_name)}</h1>
            <p class="lead">{html.escape(company_slogan)}</p>
            <div class="hero-actions">
              <a class="primary-btn" href="#quote">Request a quote</a>
              <a class="cta-btn" href="https://wa.me/27717296989?text=Hi%20Smart%20Move%20Co%2C%20I%20want%20to%20speak%20to%20an%20advisor." target="_blank" rel="noreferrer">Speak to an advisor</a>
            </div>
          </div>
          <div class="surface hero-panel">
            <h3>Why businesses choose us</h3>
            <div class="stat-grid">
              <div class="stat"><strong>Quick</strong><span>Fast quotation turnaround</span></div>
              <div class="stat"><strong>Clear</strong><span>Friendly expert guidance</span></div>
              <div class="stat"><strong>Trusted</strong><span>Professional support</span></div>
            </div>
            <div style="margin-top: 18px; color: var(--muted); line-height: 1.8;">
              We help individuals and businesses with registrations, compliance, tax, branding, and ongoing business support.
            </div>
          </div>
        </div>
      </section>

      <section id="services" class="container feature-grid">
        <div class="section-header">
          <div>
            <div class="section-label">Popular services</div>
            <h2>Solutions designed around the way you do business</h2>
          </div>
        </div>
        <div class="mini-card-grid">
          {hero_cards}
        </div>
      </section>

      <section id="quote" class="container form-shell">
        <div class="form-wrap">
          <div class="info-panel">
            <h3>Request your quotation</h3>
            <p>Share your details and select the service you need. Our team will respond with the relevant quotation and next steps.</p>
            <ul class="bullet-list">
              <li>Professional business support</li>
              <li>Fast response times</li>
              <li>Clear service guidance</li>
              <li>Custom quotes for your needs</li>
            </ul>
          </div>

          <form id="quoteForm">
            <div class="field">
              <label for="fullName">Full name</label>
              <input id="fullName" name="fullName" type="text" placeholder="Enter your full name" required />
            </div>

            <div class="field">
              <label for="contactNumber">Contact number</label>
              <input id="contactNumber" name="contactNumber" type="tel" placeholder="Enter your phone number" required />
            </div>

            <div class="field">
              <label for="email">Email address</label>
              <input id="email" name="email" type="email" placeholder="Enter your email address" required />
            </div>

            <div class="field">
              <label for="service">Service required</label>
              <select id="service" name="service" required>
                <option value="">Select a service</option>
                {service_options}
              </select>
            </div>

            <div class="submit-row">
              <button type="submit" class="primary-btn">Continue to quotation form</button>
            </div>
          </form>
        </div>
      </section>
    </main>

    <footer class="container foot">
      <div>© {html.escape(company_name)} {html.escape(founded_year)} · {html.escape(website_domain)}</div>
    </footer>

    <script>
      const quoteTarget = "#quote";
      document.getElementById('quoteForm').addEventListener('submit', function (event) {{
        event.preventDefault();
        const fullName = document.getElementById('fullName').value.trim();
        const phone = document.getElementById('contactNumber').value.trim();
        const email = document.getElementById('email').value.trim();
        const service = document.getElementById('service').value.trim();
        const params = new URLSearchParams();
        if (fullName) params.set('fullName', fullName);
        if (phone) params.set('contactNumber', phone);
        if (email) params.set('email', email);
        if (service) params.set('service', service);

        const redirectUrl = new URL(quoteTarget, window.location.href);
        for (const [key, value] of params.entries()) {{
          redirectUrl.searchParams.set(key, value);
        }}
        window.history.replaceState({{}}, '', redirectUrl.toString());
        window.scrollTo({{ top: 0, behavior: 'smooth' }});
      }});
    </script>
  </body>
</html>
"""

OUT.write_text(page, encoding="utf-8")
print(f"Generated {OUT.name} with {len(services)} active services and {len(featured_services)} featured cards.")
print(f"Company: {company_name}")
print(f"Website: {website_domain}")
print(f"Quote target: {request_form_url}")
